import io
import json
import unittest
import zipfile

import fitz
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops
from reportlab.pdfgen import canvas

from app.sandbox_artifacts import (
    COMPLIANCE_VERSION,
    SANDBOX_WATERMARK,
    build_project_application_draft_pdf,
    build_trial_ledger_xlsx,
    stamp_csv_bytes,
    stamp_json_bytes,
    stamp_pdf_bytes,
    stamp_png_bytes,
    stamp_xlsx_bytes,
    stamp_zip_bytes,
)


SOURCES = ["已发布试验资料包：R9-TEST-001"]
MODEL = "test-provider/test-model-v9"


class SandboxArtifactTests(unittest.TestCase):
    def test_pdf_has_visible_stamp_and_machine_readable_provenance(self):
        raw = io.BytesIO()
        pdf = canvas.Canvas(raw)
        pdf.drawString(72, 720, "R9 test")
        pdf.save()
        content = stamp_pdf_bytes(raw.getvalue(), sources=SOURCES, model_version=MODEL)
        document = fitz.open(stream=content, filetype="pdf")
        self.assertIn(SANDBOX_WATERMARK, document[0].get_text())
        self.assertIn("LONGYUN-AGENT-SANDBOX", document[0].get_text())
        self.assertIn(f"watermark={SANDBOX_WATERMARK}", document.metadata["keywords"])
        self.assertIn("data_sources=", document.metadata["keywords"])
        self.assertIn(f"model_version={MODEL}", document.metadata["keywords"])
        document.close()

    def test_csv_has_required_columns(self):
        content = stamp_csv_bytes("name,value\nA,1\n".encode(), sources=SOURCES, model_version=MODEL)
        text = content.decode("utf-8-sig")
        self.assertIn(SANDBOX_WATERMARK, text)
        self.assertIn(SOURCES[0], text)
        self.assertIn(MODEL, text)

    def test_json_has_compliance_object(self):
        content = stamp_json_bytes(b'{"ok": true}', sources=SOURCES, model_version=MODEL)
        payload = json.loads(content)
        self.assertEqual(payload["_sandbox_compliance"]["watermark"], SANDBOX_WATERMARK)
        self.assertEqual(payload["_sandbox_compliance"]["model_version"], MODEL)

    def test_xlsx_has_visible_compliance_sheet(self):
        workbook = Workbook()
        workbook.active.append(["name", "value"])
        workbook.active.append(["A", 1])
        raw = io.BytesIO()
        workbook.save(raw)
        content = stamp_xlsx_bytes(raw.getvalue(), sources=SOURCES, model_version=MODEL)
        stamped = load_workbook(io.BytesIO(content))
        self.assertIn("沙盒说明", stamped.sheetnames)
        values = dict(stamped["沙盒说明"].iter_rows(values_only=True))
        self.assertEqual(values["导出标识"], SANDBOX_WATERMARK)
        self.assertEqual(values["模型版本"], MODEL)

    def test_png_has_visible_pixels_and_metadata(self):
        raw = io.BytesIO()
        Image.new("RGB", (900, 500), "white").save(raw, "PNG")
        content = stamp_png_bytes(raw.getvalue(), sources=SOURCES, model_version=MODEL)
        stamped = Image.open(io.BytesIO(content))
        self.assertEqual(stamped.info["longyun_sandbox"], SANDBOX_WATERMARK)
        self.assertEqual(stamped.info["model_version"], MODEL)
        self.assertIsNotNone(ImageChops.difference(stamped.convert("RGB"), Image.new("RGB", stamped.size, "white")).getbbox())

    def test_zip_contains_compliance_manifest(self):
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, "w") as archive:
            archive.writestr("result.csv", "name,value\nA,1\n")
        content = stamp_zip_bytes(raw.getvalue(), sources=SOURCES, model_version=MODEL)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            manifest = archive.read("隆耘Agent沙盒说明.txt").decode("utf-8")
            child = archive.read("result.csv").decode("utf-8-sig")
        self.assertIn(SANDBOX_WATERMARK, manifest)
        self.assertIn(MODEL, manifest)
        self.assertIn(SANDBOX_WATERMARK, child)

    def test_new_trial_ledger_and_application_draft_are_compliant(self):
        run = {
            "id": "run-r9",
            "package_code": "R9-TEST-001",
            "package_name": "R9 验收资料包",
            "request_question": "比较不同材料的多年多点表现",
            "completed_at": "2026-09-14T10:00:00+08:00",
            "result_json": {
                "analysis_type": "multi_environment_stability",
                "analysis_version": "test-v1",
                "source_record_count": 24,
                "material_stability": [{"material_name": "测试材料A", "mean_yield_kg_per_mu": 612.5}],
            },
        }
        ledger = build_trial_ledger_xlsx(run, sources=SOURCES, model_version=MODEL)
        values = dict(load_workbook(io.BytesIO(ledger))["沙盒说明"].iter_rows(values_only=True))
        self.assertEqual(values["导出标识"], SANDBOX_WATERMARK)
        draft = build_project_application_draft_pdf(
            run, project_name="海南南繁水稻研究", sources=SOURCES, model_version=MODEL,
        )
        document = fitz.open(stream=draft, filetype="pdf")
        self.assertIn(SANDBOX_WATERMARK, document[0].get_text())
        self.assertIn(COMPLIANCE_VERSION, document.metadata["keywords"])
        document.close()


if __name__ == "__main__":
    unittest.main()
