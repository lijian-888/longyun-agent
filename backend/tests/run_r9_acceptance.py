"""Generate and verify ten cross-function artifacts for AC9.1-AC9.4."""

from __future__ import annotations

import io
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import fitz
from openpyxl import load_workbook

from app.breeding_intelligence import build_intelligence_pdf
from app.research_report import build_research_report_pdf
from app.sandbox_artifacts import (
    COMPLIANCE_VERSION,
    SANDBOX_WATERMARK,
    build_project_application_draft_pdf,
    build_trial_ledger_xlsx,
)


MODEL = "acceptance/test-model-v9"
SOURCE = "已发布试验资料包：AC9-DEMO-001"


def material_result(index: int) -> dict:
    return {
        "analysis_version": "germplasm-v1",
        "summary": f"测试材料 {index} 的可追溯种质解析。",
        "sections": {"basic": {"available": True, "records": [{"material_key": f"M{index:03d}", "name": f"测试材料{index}"}]}},
        "uncertainties": ["仅用于沙盒验收"],
        "sources": [{"batch_id": "AC9-DEMO-001", "dataset_type": "germplasm", "file_name": "germplasm.csv", "entity_type": "material", "entity_key": f"M{index:03d}"}],
    }


def parent_result(index: int) -> dict:
    source = {"batch_id": "AC9-DEMO-001", "dataset_type": "pedigree", "file_name": "pedigree.csv", "entity_type": "material", "entity_key": f"P{index:03d}"}
    return {
        "analysis_version": "parent-v1",
        "disclaimer": "亲本组合仅作辅助推荐，须由育种人员复核。",
        "breeding_goal": "高产稳产",
        "method": "受控规则评分",
        "recommendations": [{
            "female_parent": {"material_key": f"F{index}", "name": f"母本{index}"},
            "male_parent": {"material_key": f"P{index}", "name": f"父本{index}"},
            "score": 86,
            "confidence": "中",
            "recommendation_reasons": ["性状互补"],
            "risks": ["需田间验证"],
            "data_gaps": ["缺少下一季数据"],
            "evidence": [{"source": source}],
        }],
    }


def trial_analysis(index: int) -> dict:
    return {
        "analysis_type": "multi_environment_stability",
        "analysis_version": "trial-v1",
        "title": f"多年多点稳产分析 {index}",
        "model_formula": "环境内均值与变异系数",
        "source_record_count": 24,
        "material_stability": [{"material_name": f"测试材料{index}", "mean_yield_kg_per_mu": 600 + index, "cv_percent": 8.1}],
        "limitations": "仅覆盖验收资料包。",
    }


def validate_pdf(path: Path) -> None:
    document = fitz.open(path)
    try:
        keywords = document.metadata.get("keywords") or ""
        assert f"watermark={SANDBOX_WATERMARK}" in keywords
        assert "data_sources=" in keywords and "model_version=" in keywords
        for page in document:
            text = page.get_text()
            assert SANDBOX_WATERMARK in text
            assert "LONGYUN-AGENT-SANDBOX" in text
    finally:
        document.close()


def validate_xlsx(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        values = dict(workbook["沙盒说明"].iter_rows(values_only=True))
        assert values["导出标识"] == SANDBOX_WATERMARK
        assert values["数据来源"] and values["模型版本"]
        assert values["合规版本"] == COMPLIANCE_VERSION
    finally:
        workbook.close()


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="longyun-r9-") as folder:
        target = Path(folder)
        artifacts: list[tuple[str, Path]] = []
        for index in (1, 2):
            path = target / f"{index:02d}-germplasm-analysis.pdf"
            path.write_bytes(build_intelligence_pdf("种质资源综合解析报告", material_result(index)))
            artifacts.append(("种质解析报告", path))
        for index in (1, 2):
            path = target / f"{index + 2:02d}-parent-recommendation.pdf"
            path.write_bytes(build_intelligence_pdf("亲本组合辅助推荐报告", parent_result(index)))
            artifacts.append(("亲本辅助推荐报告", path))
        for index in (1, 2):
            analysis = trial_analysis(index)
            evidence = [{"priority": "P0", "title": SOURCE, "detail": "24 条已发布记录"}]
            path = target / f"{index + 4:02d}-trial-analysis.pdf"
            path.write_bytes(build_research_report_pdf(question="分析多年多点表现", answer="已完成受控统计分析。", evidence=evidence, analysis=analysis))
            artifacts.append(("试验分析报告", path))
        for index in (1, 2):
            analysis = trial_analysis(index)
            run = {"id": f"run-{index}", "package_code": "AC9-DEMO-001", "package_name": "AC9 验收资料包", "request_question": "分析多年多点表现", "result_json": analysis}
            path = target / f"{index + 6:02d}-trial-ledger.xlsx"
            path.write_bytes(build_trial_ledger_xlsx(run, sources=[SOURCE], model_version=MODEL))
            artifacts.append(("试验台账", path))
        for index in (1, 2):
            analysis = trial_analysis(index)
            run = {"id": f"run-{index}", "package_code": "AC9-DEMO-001", "package_name": "AC9 验收资料包", "request_question": "分析多年多点表现", "result_json": analysis}
            path = target / f"{index + 8:02d}-project-application.pdf"
            path.write_bytes(build_project_application_draft_pdf(run, project_name="海南南繁水稻研究", sources=[SOURCE], model_version=MODEL))
            artifacts.append(("课题申报辅助材料草稿", path))
        for _, path in artifacts:
            validate_xlsx(path) if path.suffix == ".xlsx" else validate_pdf(path)
        categories = sorted({category for category, _ in artifacts})
        assert len(artifacts) == 10 and len(categories) == 5
        print(json.dumps({"status": "passed", "verified": len(artifacts), "categories": categories, "watermark": SANDBOX_WATERMARK}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
