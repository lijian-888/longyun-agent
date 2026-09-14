"""Uniform compliance controls for every user-downloadable sandbox artifact.

The source files used by local analysis remain unchanged.  Only exported copies
are decorated, so watermarks cannot influence subsequent statistics or uploads.
"""

from __future__ import annotations

import csv
import html
import io
import json
import os
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


SANDBOX_WATERMARK = "隆耘 Agent 沙盒演示环境"
COMPLIANCE_VERSION = "r9-sandbox-artifact-v1"
_PDF_MARKER = "LONGYUN-AGENT-SANDBOX"


def artifact_model_version(analysis: dict[str, Any] | None = None) -> str:
    """Return a non-secret, server-derived model/engine version label."""
    provider = (os.getenv("AI_PROVIDER") or "shennong").strip().lower()
    if provider == "vllm":
        model = (os.getenv("VLLM_MODEL") or "local-model").strip()
    else:
        provider = "shennong"
        model = (os.getenv("SHENNONG_MODEL") or "sn").strip()
    parts = [f"{provider}/{model or 'unknown'}"]
    if analysis:
        engine = analysis.get("engine") or analysis.get("analysis_engine")
        version = analysis.get("analysis_version") or analysis.get("rule_version")
        if engine or version:
            parts.append(f"{engine or 'controlled-engine'}/{version or 'version-not-recorded'}")
    return "；".join(parts)


def source_descriptions(sources: Iterable[Any] | None) -> list[str]:
    values: list[str] = []
    for source in sources or []:
        if isinstance(source, str):
            value = source.strip()
        elif isinstance(source, dict):
            title = source.get("title") or source.get("source") or source.get("file_name") or source.get("dataset_type")
            locator = source.get("url") or source.get("source_locator") or source.get("batch_id") or source.get("detail")
            value = "：".join(str(item).strip() for item in (title, locator) if item)
        else:
            value = str(source).strip()
        if value and value not in values:
            values.append(value[:500])
    return values[:20] or ["当前账号有权访问的已发布标准数据或本轮可追溯分析结果"]


def compliance_metadata(
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> dict[str, Any]:
    return {
        "watermark": SANDBOX_WATERMARK,
        "data_sources": source_descriptions(sources),
        "model_version": model_version or artifact_model_version(),
        "compliance_version": COMPLIANCE_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    }


def _pdf_font_file() -> str | None:
    candidates = (
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    )
    return str(next((path for path in candidates if path.is_file()), "")) or None


def stamp_pdf_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    document = fitz.open(stream=content, filetype="pdf")
    existing = document.metadata or {}
    if _PDF_MARKER in str(existing.get("keywords") or ""):
        document.close()
        return content
    for page in document:
        page_rect = page.rect
        # MuPDF's built-in simplified-Chinese font keeps the stamp portable and
        # avoids embedding a 10+ MB system TTC in every exported page.
        font_name = "china-s"
        watermark_box = fitz.Rect(
            page_rect.x0 + 16,
            page_rect.y0 + page_rect.height * 0.43,
            page_rect.x1 - 16,
            page_rect.y0 + page_rect.height * 0.57,
        )
        page.insert_textbox(
            watermark_box,
            SANDBOX_WATERMARK,
            fontname=font_name,
            fontsize=max(20, min(34, page_rect.width / 16)),
            color=(0.58, 0.62, 0.60),
            fill_opacity=0.18,
            align=fitz.TEXT_ALIGN_CENTER,
            overlay=True,
        )
        footer = (
            f"数据来源：{'；'.join(metadata['data_sources'][:2])} | "
            f"模型版本：{metadata['model_version']}"
        )[:180]
        page.insert_textbox(
            fitz.Rect(page_rect.x0 + 18, page_rect.y1 - 25, page_rect.x1 - 18, page_rect.y1 - 11),
            footer,
            fontname=font_name,
            fontsize=5.5,
            color=(0.34, 0.42, 0.39),
            align=fitz.TEXT_ALIGN_CENTER,
            overlay=True,
        )
        page.insert_text(
            (page_rect.x0 + 18, page_rect.y1 - 5),
            _PDF_MARKER,
            fontname="helv",
            fontsize=4.5,
            color=(0.45, 0.50, 0.48),
            overlay=True,
        )
    keywords = "; ".join((
        _PDF_MARKER,
        f"watermark={SANDBOX_WATERMARK}",
        f"data_sources={' | '.join(metadata['data_sources'])}",
        f"model_version={metadata['model_version']}",
        f"compliance_version={COMPLIANCE_VERSION}",
    ))
    document.set_metadata({
        **existing,
        "producer": "Longyun Agent sandbox artifact service",
        "keywords": keywords,
    })
    result = document.tobytes(garbage=3, deflate=True)
    document.close()
    return result


def stamp_csv_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
    delimiter: str = ",",
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    text = content.decode("utf-8-sig", errors="replace")
    if SANDBOX_WATERMARK in text:
        return content
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        rows = [["导出内容"]]
    rows[0].extend(["沙盒水印", "数据来源", "模型版本"])
    if len(rows) == 1:
        rows.append([""] * max(1, len(rows[0]) - 3))
    rows[1].extend([
        SANDBOX_WATERMARK,
        "；".join(metadata["data_sources"]),
        metadata["model_version"],
    ])
    output = io.StringIO()
    writer = csv.writer(output, delimiter=delimiter, lineterminator="\r\n")
    writer.writerows(rows)
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def stamp_json_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        payload = {"raw_text": content.decode("utf-8", errors="replace")}
    if isinstance(payload, dict) and payload.get("_sandbox_compliance", {}).get("watermark") == SANDBOX_WATERMARK:
        return content
    if not isinstance(payload, dict):
        payload = {"records": payload}
    payload["_sandbox_compliance"] = compliance_metadata(sources, model_version)
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def stamp_xlsx_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    workbook = load_workbook(io.BytesIO(content))
    if "沙盒说明" in workbook.sheetnames:
        sheet = workbook["沙盒说明"]
    else:
        sheet = workbook.create_sheet("沙盒说明")
    sheet.delete_rows(1, sheet.max_row)
    rows = [
        ("导出标识", SANDBOX_WATERMARK),
        ("数据来源", "；".join(metadata["data_sources"])),
        ("模型版本", metadata["model_version"]),
        ("合规版本", COMPLIANCE_VERSION),
        ("生成时间", metadata["generated_at"]),
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 100
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["B1"].font = Font(bold=True, color="FFFFFF", size=16)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="2E6F5E")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for data_sheet in workbook.worksheets:
        data_sheet.oddHeader.center.text = SANDBOX_WATERMARK
        data_sheet.oddFooter.center.text = f"数据来源：{'；'.join(metadata['data_sources'][:2])} | 模型版本：{metadata['model_version']}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _image_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    font_file = _pdf_font_file()
    if font_file:
        return ImageFont.truetype(font_file, size, index=0)
    return ImageFont.load_default()


def stamp_png_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    source_image = Image.open(io.BytesIO(content))
    if source_image.info.get("longyun_sandbox") == SANDBOX_WATERMARK:
        return content
    image = source_image.convert("RGBA")
    overlay = Image.new("RGBA", image.size, (255, 255, 255, 0))
    draw = ImageDraw.Draw(overlay)
    watermark_font = _image_font(max(18, min(44, image.width // 20)))
    footer_font = _image_font(max(10, min(18, image.width // 65)))
    box = draw.textbbox((0, 0), SANDBOX_WATERMARK, font=watermark_font)
    text_width = box[2] - box[0]
    draw.text(
        ((image.width - text_width) / 2, image.height * 0.46),
        SANDBOX_WATERMARK,
        font=watermark_font,
        fill=(80, 105, 96, 62),
    )
    footer_height = max(34, image.height // 16)
    draw.rectangle((0, image.height - footer_height, image.width, image.height), fill=(255, 255, 255, 220))
    footer = f"{SANDBOX_WATERMARK}｜数据来源：{'；'.join(metadata['data_sources'][:2])}｜模型版本：{metadata['model_version']}"
    draw.text((12, image.height - footer_height + 8), footer[:180], font=footer_font, fill=(44, 78, 67, 255))
    result = Image.alpha_composite(image, overlay).convert("RGB")
    output = io.BytesIO()
    png_info = PngImagePlugin.PngInfo()
    png_info.add_text("longyun_sandbox", SANDBOX_WATERMARK)
    png_info.add_text("data_sources", "；".join(metadata["data_sources"]))
    png_info.add_text("model_version", metadata["model_version"])
    png_info.add_text("compliance_version", COMPLIANCE_VERSION)
    result.save(output, format="PNG", optimize=True, pnginfo=png_info)
    return output.getvalue()


def stamp_text_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    text = content.decode("utf-8-sig", errors="replace")
    if SANDBOX_WATERMARK in text:
        return content
    prefix = (
        f"# {SANDBOX_WATERMARK}\n"
        f"# 数据来源：{'；'.join(metadata['data_sources'])}\n"
        f"# 模型版本：{metadata['model_version']}\n\n"
    )
    return (prefix + text).encode("utf-8")


def stamp_zip_bytes(
    content: bytes,
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    metadata = compliance_metadata(sources, model_version)
    source = io.BytesIO(content)
    output = io.BytesIO()
    with zipfile.ZipFile(source, "r") as archive_in, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive_out:
        if "隆耘Agent沙盒说明.txt" in archive_in.namelist():
            return content
        for item in archive_in.infolist():
            data = archive_in.read(item.filename)
            suffix = Path(item.filename).suffix.lower()
            if not item.is_dir():
                data = ensure_compliant_artifact_bytes(
                    data,
                    file_name=item.filename,
                    content_type="",
                    sources=metadata["data_sources"],
                    model_version=metadata["model_version"],
                    _inside_zip=True,
                )
            archive_out.writestr(item, data)
        manifest = (
            f"{SANDBOX_WATERMARK}\n"
            f"数据来源：{'；'.join(metadata['data_sources'])}\n"
            f"模型版本：{metadata['model_version']}\n"
            f"合规版本：{COMPLIANCE_VERSION}\n"
        )
        archive_out.writestr("隆耘Agent沙盒说明.txt", manifest.encode("utf-8"))
    return output.getvalue()


def ensure_compliant_artifact_bytes(
    content: bytes,
    *,
    file_name: str,
    content_type: str,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
    _inside_zip: bool = False,
) -> bytes:
    """Decorate a downloadable copy according to its real file format."""
    suffix = Path(file_name).suffix.lower()
    normalized_type = (content_type or "").split(";", 1)[0].strip().lower()
    if suffix == ".pdf" or normalized_type == "application/pdf":
        return stamp_pdf_bytes(content, sources=sources, model_version=model_version)
    if suffix == ".png" or normalized_type == "image/png":
        return stamp_png_bytes(content, sources=sources, model_version=model_version)
    if suffix == ".xlsx" or normalized_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return stamp_xlsx_bytes(content, sources=sources, model_version=model_version)
    if suffix == ".csv" or normalized_type == "text/csv":
        return stamp_csv_bytes(content, sources=sources, model_version=model_version)
    if suffix in {".tsv", ".txt", ".log", ".md"} or normalized_type.startswith("text/"):
        delimiter = "\t" if suffix == ".tsv" else None
        if delimiter:
            return stamp_csv_bytes(content, sources=sources, model_version=model_version, delimiter=delimiter)
        return stamp_text_bytes(content, sources=sources, model_version=model_version)
    if suffix == ".json" or normalized_type == "application/json":
        return stamp_json_bytes(content, sources=sources, model_version=model_version)
    if suffix == ".zip" or normalized_type == "application/zip":
        if _inside_zip:
            return content
        return stamp_zip_bytes(content, sources=sources, model_version=model_version)
    return content


def build_trial_ledger_xlsx(
    run: dict[str, Any],
    *,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    analysis = run.get("result_json") or {}
    model_version = model_version or artifact_model_version(analysis)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "试验台账"
    summary.append(["项目", "内容"])
    summary_rows = [
        ("成果类型", "试验台账"),
        ("资料包编号", run.get("package_code") or "-"),
        ("资料包名称", run.get("package_name") or "-"),
        ("分析运行编号", run.get("id") or run.get("analysis_run_id") or "-"),
        ("分析问题", run.get("request_question") or "-"),
        ("分析类型", analysis.get("analysis_type") or "-"),
        ("统计方法", analysis.get("model_formula") or analysis.get("method") or "-"),
        ("记录数量", analysis.get("source_record_count") or analysis.get("sample_size") or 0),
        ("完成时间", run.get("completed_at") or "-"),
        ("局限说明", analysis.get("limitations") or "当前台账仅覆盖已发布、可追溯的试验数据。"),
    ]
    for row in summary_rows:
        summary.append(row)
    details = workbook.create_sheet("分析明细")
    details.append(["结果分区", "序号", "字段", "值"])
    for section, values in analysis.items():
        if not isinstance(values, list):
            continue
        for index, record in enumerate(values, start=1):
            if isinstance(record, dict):
                for key, value in record.items():
                    details.append([section, index, key, json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else value])
            else:
                details.append([section, index, "value", str(record)])
    if details.max_row == 1:
        details.append(["summary", 1, "result", json.dumps(analysis, ensure_ascii=False, default=str)])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E6F5E")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column, width in (("A", 24), ("B", 30), ("C", 32), ("D", 80)):
            sheet.column_dimensions[column].width = width
    raw = io.BytesIO()
    workbook.save(raw)
    return stamp_xlsx_bytes(raw.getvalue(), sources=sources, model_version=model_version)


def build_project_application_draft_pdf(
    run: dict[str, Any],
    *,
    project_name: str,
    sources: Iterable[Any] | None = None,
    model_version: str | None = None,
) -> bytes:
    analysis = run.get("result_json") or {}
    model_version = model_version or artifact_model_version(analysis)
    source_values = source_descriptions(sources)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    body = styles["BodyText"].clone("ApplicationBody")
    body.fontName = "STSong-Light"
    body.fontSize = 9.5
    body.leading = 15
    heading = styles["Heading2"].clone("ApplicationHeading")
    heading.fontName = "STSong-Light"
    heading.textColor = colors.HexColor("#174D3E")
    title = styles["Title"].clone("ApplicationTitle")
    title.fontName = "STSong-Light"
    title.textColor = colors.HexColor("#174D3E")

    def paragraph(value: Any, style: Any = body) -> Paragraph:
        return Paragraph(html.escape(str(value or "-")).replace("\n", "<br/>") , style)

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=16 * mm,
        title=f"{project_name}课题申报辅助材料草稿",
    )
    record_count = analysis.get("source_record_count") or analysis.get("sample_size") or 0
    story: list[Any] = [
        paragraph("课题申报辅助材料草稿", title),
        paragraph(f"课题：{project_name}｜资料包：{run.get('package_name') or '-'}（{run.get('package_code') or '-'}）"),
        paragraph(f"{SANDBOX_WATERMARK}；本稿用于材料组织演示，不构成正式申报文本。"),
        Spacer(1, 4 * mm),
        paragraph("一、立项依据与问题", heading),
        paragraph(run.get("request_question") or "基于当前已发布试验数据形成可复核的科研问题。"),
        paragraph("二、已有数据基础", heading),
        paragraph(f"当前受控分析纳入 {record_count} 条可追溯记录，分析类型为 {analysis.get('title') or analysis.get('analysis_type') or '受控试验分析'}。"),
        paragraph("三、拟解决的科学与业务目标", heading),
        paragraph("围绕材料表现、环境稳定性和关键性状证据开展复核；未被当前数据支持的目标、指标和结论须由申报团队补充，不由系统推断。"),
        paragraph("四、研究内容与技术路线（草稿）", heading),
        paragraph(f"采用 {analysis.get('model_formula') or analysis.get('method') or '经审核的受控统计流程'} 对已发布数据进行分析，保留筛选条件、输入记录、统计版本和结果追溯信息。"),
        paragraph("五、预期成果（草稿）", heading),
        paragraph("形成可追溯的试验分析报告、试验台账和阶段性材料评价结果；正式考核指标、经费和进度需由负责人审核补齐。"),
        paragraph("六、风险、局限与人工复核", heading),
        paragraph(analysis.get("limitations") or "当前结论仅覆盖已发布数据，缺失试验、未公开数据及未审核材料不得作为已完成证据。"),
        paragraph("七、数据来源和模型版本说明", heading),
        paragraph("数据来源：" + "；".join(source_values)),
        paragraph("模型版本：" + model_version),
    ]
    document.build(story)
    return stamp_pdf_bytes(buffer.getvalue(), sources=source_values, model_version=model_version)
